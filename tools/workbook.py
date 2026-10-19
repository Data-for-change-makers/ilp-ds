"""Make and read ILP-DS review workbooks (Excel).

A workbook is how a publisher drafts and reviews records before they become an ILP-DS file:
one row per record, with a dropdown for every field that has a code list.

Usage:
    python tools/workbook.py read WORKBOOK.xlsx --out INTERNAL.geojson [--public PUBLIC.geojson]
                                 [--geometry SHAPES.geojson] [--display-fields FIELDS.json]

``make_workbook()`` is called from a publisher's own script with its records. ``read`` turns a
reviewed workbook back into an internal ILP-DS file, checks it, and can write the public copy.

Workbook layout
---------------
* Read me: how to use the workbook, and what the cell and header colours mean.
* Records: row 1 groups, row 2 labels, row 3 ILP-DS keys (do not edit), row 4 hints, data from
  row 5. Columns named ``field.1``, ``field.2`` ... are one multi-value field. ``review:``
  columns are notes for reviewers and are never converted. ``x-`` columns are the publisher's
  own fields; a Yes/No column becomes true/false.
* Lists: the code lists (label and value) behind the dropdowns.
* Fields: one row per column, with the definition and, optionally, whether a display (such as a
  map) shows the field.
* File: publisher, licence, base and a change note.

A cell's colour can record where its value came from, and its comment the page or the reason.
A value in a column that does not apply to the record's kind is kept as ``x-<field>`` and
reported, so nothing typed in the workbook is lost silently.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))

with (ROOT / "schema" / "fields.csv").open(newline="", encoding="utf-8") as _h:
    REGISTRY = {row["field"]: row for row in csv.DictReader(_h)}

FIRST_ROW = 5
LAST_ROW = 300
THIN = Side(style="thin", color="BFBFBF")

# Where a value came from: fill colour and legend text. Publishers may pass their own.
DEFAULT_SOURCES = {
    "primary": ("E2EFDA", "From the main source (page or link in the comment)"),
    "secondary": ("DDEBF7", "From another source (named in the comment)"),
    "inferred": ("FFF2CC", "Inferred or converted: please check (reason in the comment)"),
    "gap": ("F8CBAD", "A decision is needed (see the comment)"),
}
HEADER_FILL = {
    "public": "1F4E5A", "restricted": "C65911", "internal": "A50021",
    "display": "375623", "review": "7F7F7F",
}

LONDON_BOROUGHS = [
    ("City of London", "E09000001"), ("Barking and Dagenham", "E09000002"), ("Barnet", "E09000003"),
    ("Bexley", "E09000004"), ("Brent", "E09000005"), ("Bromley", "E09000006"), ("Camden", "E09000007"),
    ("Croydon", "E09000008"), ("Ealing", "E09000009"), ("Enfield", "E09000010"),
    ("Greenwich", "E09000011"), ("Hackney", "E09000012"), ("Hammersmith and Fulham", "E09000013"),
    ("Haringey", "E09000014"), ("Harrow", "E09000015"), ("Havering", "E09000016"),
    ("Hillingdon", "E09000017"), ("Hounslow", "E09000018"), ("Islington", "E09000019"),
    ("Kensington and Chelsea", "E09000020"), ("Kingston upon Thames", "E09000021"),
    ("Lambeth", "E09000022"), ("Lewisham", "E09000023"), ("Merton", "E09000024"),
    ("Newham", "E09000025"), ("Redbridge", "E09000026"), ("Richmond upon Thames", "E09000027"),
    ("Southwark", "E09000028"), ("Sutton", "E09000029"), ("Tower Hamlets", "E09000030"),
    ("Waltham Forest", "E09000031"), ("Wandsworth", "E09000032"), ("Westminster", "E09000033"),
]


def code_list(name: str) -> list[tuple[str, str]]:
    with (ROOT / "codelists" / f"{name}.csv").open(newline="", encoding="utf-8") as handle:
        return [(row["label"], row["value"]) for row in csv.DictReader(handle)]


CODE_LISTS = ("kind", "overall-status", "delivery-stage", "planning-status", "permission-type",
              "land-type", "ownership-status", "asset-class", "economic-driver", "gva-method",
              "value-type", "opportunity-type", "owner-vision", "funding-status", "land-assembly")


def col(key, label, group, width=14, lst=None, hint="", vis=None, field=None, kind="text"):
    field = field or key.split(".")[0].replace("consent:", "")
    if vis is None:
        vis = REGISTRY[field]["visibility"] if field in REGISTRY else "review"
    return {"key": key, "label": label, "group": group, "width": width, "list": lst,
            "hint": hint, "vis": vis, "field": field, "kind": kind}


def columns(display: dict | None = None) -> list[dict]:
    """The Records sheet's columns. ``display`` adds a Yes/No column such as
    {"field": "x-discover", "label": "Show on Discover", "hint": "Yes: a pin on the map"}."""
    c = []
    g = "Record"
    c += [col("id", "Id", g, 22, hint="lower-case-with-hyphens; unique", field="id", vis="public"),
          col("name", "Name", g, 30)]
    if display:
        c.append(col(display["field"], display["label"], g, 11, "yes-no", display.get("hint", ""),
                     vis="display", field=display["field"]))
    c += [col("kind", "Record type", g, 16, "kind"),
          col("parent", "Parent id", g, 20, hint="id of the record one level up")]
    c += [col(f"serves.{i}", f"Serves {i}", g, 16, hint="infrastructure only: ids it serves")
          for i in range(1, 7)]
    c += [col("description", "Description", g, 60, hint="up to 400 words")]
    g = "Place"
    c += [col("organisation", "Lead body", g, 30)]
    c += [col(f"local-authorities.{i}", f"Local authority {i}", g, 16, "local-authority")
          for i in range(1, 5)]
    c += [col("planning-authority", "Planning authority (if not the council)", g, 22,
              hint="Planning Data reference, e.g. development-corporation:Q20648596"),
          col("opportunity-area", "Opportunity Area", g, 20),
          col("designated-area", "In a designated area?", g, 11, "yes-no"),
          col("designations", "Designations", g, 24, hint="separate with ;"),
          col("land-type", "Land type", g, 16, "land-type"),
          col("ownership-status", "Ownership", g, 18, "ownership-status"),
          col("hectares", "Hectares", g, 10, hint="number", kind="number"),
          col("point.lon", "Longitude", g, 10, hint="WGS84", field="point", kind="number"),
          col("point.lat", "Latitude", g, 10, hint="WGS84", field="point", kind="number"),
          col("review:geometry", "Map shape", g, 26, hint="boundary or line used, if any")]
    g = "Status and planning"
    c += [col("overall-status", "Overall status", g, 14, "overall-status"),
          col("delivery-stage", "Delivery stage", g, 15, "delivery-stage",
              "leave blank for a whole area"),
          col("planning-status", "Planning status", g, 15, "planning-status"),
          col("permission-type", "Permission type", g, 18, "permission-type"),
          col("permission-date", "Permission date", g, 12, hint="YYYY-MM-DD", kind="date"),
          col("planning-references", "Planning references", g, 18, hint="separate with ;"),
          col("planning-documents-url", "Planning documents link", g, 24)]
    g = "Scale"
    c += [col("homes-consented", "Homes consented", g, 11, kind="integer"),
          col("homes-pipeline", "Homes in vision, beyond consent", g, 13, kind="integer",
              hint="never include consented homes"),
          col("floorspace-consented.total", "Floorspace consented (m²)", g, 13, kind="integer"),
          col("floorspace-pipeline.total", "Floorspace in vision, beyond consent (m²)", g, 14,
              kind="integer", hint="single figure, or the lower figure of a range"),
          col("floorspace-pipeline-to.total", "Floorspace in vision: upper figure of range (m²)", g, 14,
              kind="integer"),
          col("jobs-from", "Jobs", g, 10, hint="single figure, or the lower figure of a range",
              kind="integer"),
          col("jobs-to", "Jobs: upper figure of range", g, 11, kind="integer")]
    g = "Drivers"
    c += [col(f"asset-classes.{i}", f"Asset class {i}", g, 18, "asset-class") for i in range(1, 6)]
    c += [col(f"economic-drivers.{i}", f"Sector {i}", g, 18, "economic-driver") for i in range(1, 6)]
    g = "Estimates"
    c += [col("gva-uplift", "GVA uplift (£ a year)", g, 15, kind="number"),
          col("gva-method", "GVA method", g, 12, "gva-method"),
          col("gva-source", "GVA source", g, 30),
          col("gva-date", "GVA date", g, 12, hint="YYYY-MM-DD", kind="date"),
          col("investment-value", "Investment figure (£)", g, 16, kind="number"),
          col("value-type", "What the figure measures", g, 18, "value-type"),
          col("value-source", "Figure source (quote the wording)", g, 34),
          col("value-date", "Figure date", g, 12, hint="YYYY-MM-DD", kind="date")]
    g = "How to take part"
    c += [col(f"opportunity-types.{i}", f"Opportunity type {i}", g, 18, "opportunity-type")
          for i in range(1, 6)]
    c += [col("opportunity-type-notes", "Opportunity type notes", g, 34,
              hint="detail, or describe 'Other'")]
    g = "Contacts and partners"
    for i in range(1, 4):
        c += [col(f"contacts.{i}.name", f"Contact {i} name", g, 18, hint="published if supplied"),
              col(f"contacts.{i}.role", f"Contact {i} role", g, 26),
              col(f"contacts.{i}.email", f"Contact {i} email", g, 26),
              col(f"contacts.{i}.url", f"Contact {i} web page", g, 20)]
    c += [col("investors", "Investors", g, 30, hint="separate with ;"),
          col("model-url", "3D model link", g, 20),
          col("model-commissioner", "3D model paid for by", g, 18)]
    g = "Restricted and internal"
    c += [col("owner-vision", "Owner's vision", g, 13, "owner-vision"),
          col("consent:owner-vision", "Owner agreed to publish vision?", g, 13, "yes-no", vis="restricted"),
          col("funding-status", "Funding status", g, 18, "funding-status"),
          col("consent:funding-status", "Owner agreed to publish funding?", g, 13, "yes-no",
              vis="restricted"),
          col("land-assembly", "Land assembly (never published)", g, 16, "land-assembly")]
    g = "Record keeping"
    c += [col("source", "Sources", g, 40),
          col("entry-date", "Last checked", g, 12, hint="YYYY-MM-DD", kind="date"),
          col("start-date", "Start date", g, 12, kind="date"),
          col("end-date", "End date", g, 12, kind="date")]
    g = "Review only (not published)"
    c += [col("review:status", "Status as published", g, 24),
          col("review:drivers", "Drivers as published", g, 36),
          col("review:opportunity-type", "Opportunity type as published", g, 30),
          col("review:pipeline", "Pipeline as published", g, 36),
          col("review:notes", "Notes for reviewer", g, 40)]
    return c


def make_workbook(records: list[dict], path: Path, *, file_info: dict, readme: list[str],
                  display: dict | None = None, display_fields: set[str] = frozenset(),
                  local_authorities: list[tuple[str, str]] = LONDON_BOROUGHS,
                  sources: dict = DEFAULT_SOURCES, author: str = "ILP-DS workbook") -> None:
    """Write a review workbook.

    Each record is {column key: value}, where a value may be (value, source, comment) to colour
    the cell and attach a comment. Values for list columns are given as code-list values.
    ``display`` adds a Yes/No column (see ``columns``); ``display_fields`` are the fields that
    display shows, which the Fields sheet lets the reviewer change.
    """
    cols = columns(display)
    lists = {name: code_list(name) for name in CODE_LISTS}
    lists["local-authority"] = list(local_authorities)
    lists["yes-no"] = [("Yes", "yes"), ("No", "no")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Records"

    # Lists sheet
    ls = wb.create_sheet("Lists")
    ranges = {}
    c_idx = 1
    for name, items in lists.items():
        ls.cell(1, c_idx, name).font = Font(bold=True)
        ls.cell(1, c_idx + 1, f"{name} (value)").font = Font(bold=True, color="7F7F7F")
        for r, (label, value) in enumerate(items, start=2):
            ls.cell(r, c_idx, label)
            ls.cell(r, c_idx + 1, value)
        letter = get_column_letter(c_idx)
        ranges[name] = f"Lists!${letter}$2:${letter}${len(items) + 1}"
        ls.column_dimensions[letter].width = max(14, max(len(lbl) for lbl, _ in items) + 2)
        ls.column_dimensions[get_column_letter(c_idx + 1)].width = 14
        c_idx += 3
    label_of = {name: {v: lbl for lbl, v in items} for name, items in lists.items()}

    # Header rows
    groups = []
    for i, cdef in enumerate(cols, start=1):
        if not groups or groups[-1][0] != cdef["group"]:
            groups.append([cdef["group"], i, i])
        else:
            groups[-1][2] = i
    for name, start, end in groups:
        cell = ws.cell(1, start, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="404040")
        if end > start:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
    for i, cdef in enumerate(cols, start=1):
        h = ws.cell(2, i, cdef["label"])
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=HEADER_FILL[cdef["vis"]])
        h.alignment = Alignment(wrap_text=True, vertical="top")
        k = ws.cell(3, i, cdef["key"])
        k.font = Font(size=8, color="7F7F7F")
        hint = ws.cell(4, i, cdef["hint"] or ("choose from list" if cdef["list"] else ""))
        hint.font = Font(size=8, italic=True, color="595959")
        hint.alignment = Alignment(wrap_text=True, vertical="top")
        field = cdef["field"]
        if field in REGISTRY:
            h.comment = Comment(f"{field} ({REGISTRY[field]['visibility']}): {REGISTRY[field]['definition']}",
                                "ILP-DS")
        ws.column_dimensions[get_column_letter(i)].width = cdef["width"]
        if cdef["list"]:
            dv = DataValidation(type="list", formula1=f"={ranges[cdef['list']]}", allow_blank=True,
                                showErrorMessage=True, errorTitle="Not in the list",
                                error="Choose a value from the list (see the Lists sheet).")
            ws.add_data_validation(dv)
            letter = get_column_letter(i)
            dv.add(f"{letter}{FIRST_ROW}:{letter}{LAST_ROW}")
    ws.row_dimensions[2].height = 45
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = ws.cell(FIRST_ROW, 3)

    # Data rows
    for r, rec in enumerate(records, start=FIRST_ROW):
        ws.row_dimensions[r].height = 48
        for i, cdef in enumerate(cols, start=1):
            raw = rec.get(cdef["key"])
            if raw is None:
                continue
            value, source, note = (raw + (None, None))[:3] if isinstance(raw, tuple) else (raw, None, None)
            if cdef["list"] and value is not None and value != "":
                value = label_of[cdef["list"]].get(value, value)
            if cdef["kind"] == "date" and isinstance(value, str) and re.match(r"^\d{4}-\d{2}-\d{2}$", value):
                value = dt.date.fromisoformat(value)
            cell = ws.cell(r, i, value)
            cell.alignment = Alignment(wrap_text=cdef["width"] >= 30, vertical="top")
            cell.border = Border(bottom=THIN)
            if cdef["kind"] == "date":
                cell.number_format = "yyyy-mm-dd"
            elif cdef["kind"] in ("integer", "number") and isinstance(value, (int, float)):
                cell.number_format = "#,##0.##" if cdef["kind"] == "number" else "#,##0"
                if cdef["key"] in ("point.lon", "point.lat"):
                    cell.number_format = "0.0000"
            if source:
                cell.fill = PatternFill("solid", fgColor=sources[source][0])
            if note:
                cell.comment = Comment(note, author)
                cell.comment.width, cell.comment.height = 320, 140
    ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{FIRST_ROW + max(len(records), 1) - 1}"

    # Fields sheet
    fs = wb.create_sheet("Fields")
    head = ["Column", "Records column key", "ILP-DS field", "Applies to", "Visibility", "Definition"]
    if display:
        head.append(f"{display['label']}: show this field")
    for i, h in enumerate(head, start=1):
        cell = fs.cell(1, i, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E5A")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1=f"={ranges['yes-no']}", allow_blank=True)
    fs.add_data_validation(dv)
    for r, cdef in enumerate(cols, start=2):
        reg = REGISTRY.get(cdef["field"], {})
        fs.cell(r, 1, cdef["label"])
        fs.cell(r, 2, cdef["key"]).font = Font(size=8, color="7F7F7F")
        fs.cell(r, 3, cdef["field"] if reg else "")
        fs.cell(r, 4, reg.get("applies-to", ""))
        fs.cell(r, 5, cdef["vis"])
        if reg:
            definition = reg["definition"]
        elif display and cdef["field"] == display["field"]:
            definition = display.get("definition", "The publisher's own Yes/No field.")
        else:
            definition = "Notes for reviewers. Not converted or published."
        fs.cell(r, 6, definition)
        if display:
            if cdef["vis"] == "public" and reg:
                fs.cell(r, 7, "Yes" if cdef["field"] in display_fields else "No")
                dv.add(f"G{r}")
            else:
                fs.cell(r, 7, "n/a").font = Font(color="7F7F7F")
    for letter, width in zip("ABCDEFG", (30, 24, 22, 30, 11, 70, 14)):
        fs.column_dimensions[letter].width = width
    for row in fs.iter_rows(min_row=2, max_col=6):
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    fs.freeze_panes = "B2"

    # File sheet
    fi = wb.create_sheet("File")
    fi.column_dimensions["A"].width = 18
    fi.column_dimensions["B"].width = 80
    for r, (k, v) in enumerate(file_info.items(), start=1):
        fi.cell(r, 1, k).font = Font(bold=True)
        fi.cell(r, 2, v).alignment = Alignment(wrap_text=True)

    # Read me sheet, first
    rm = wb.create_sheet("Read me", 0)
    rm.column_dimensions["A"].width = 4
    rm.column_dimensions["B"].width = 110
    r = 1
    for line in readme:
        cell = rm.cell(r, 2, line)
        if line.startswith("# "):
            cell.value = line[2:]
            cell.font = Font(bold=True, size=14)
        elif line.startswith("## "):
            cell.value = line[3:]
            cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
    rm.cell(r, 2, "Cell colours").font = Font(bold=True, size=11)
    for colour, text in sources.values():
        r += 1
        rm.cell(r, 1).fill = PatternFill("solid", fgColor=colour)
        rm.cell(r, 2, text)
    r += 1
    rm.cell(r, 2, "White: no value found; fill in if you know it, or leave blank.")
    r += 2
    rm.cell(r, 2, "Column header colours").font = Font(bold=True, size=11)
    legend = [("public", "Public: published in the open feed"),
              ("restricted", "Restricted: published only if the owner agrees (consent column)"),
              ("internal", "Internal: never published")]
    if display:
        legend.append(("display", f"{display['label']}: the publisher's own display choice"))
    legend.append(("review", "Review only: notes, never published"))
    for vis, text in legend:
        r += 1
        rm.cell(r, 1).fill = PatternFill("solid", fgColor=HEADER_FILL[vis])
        rm.cell(r, 2, text)
    wb.save(path)


# Reading a reviewed workbook ---------------------------------------------------------

INT_FIELDS = {"homes-consented", "homes-pipeline", "jobs-from", "jobs-to", "storeys"}
NUM_FIELDS = {"hectares", "gva-uplift", "investment-value", "height-m"}
DATE_FIELDS = {"permission-date", "gva-date", "value-date", "entry-date", "start-date", "end-date"}
SPLIT_FIELDS = {"designations", "planning-references", "story-stops"}


def _clean(value):
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _ref(value: str) -> str:
    return value if value.startswith("#") else "#" + value


def _version() -> str:
    text = (ROOT / "tools" / "build_schema.py").read_text(encoding="utf-8")
    return re.search(r'VERSION = "([^"]+)"', text).group(1)


def read_workbook(path: Path, geometry: dict | None = None) -> tuple[dict, list[str], dict]:
    """Return (internal ILP-DS document, notes on what was changed, display field choices)."""
    wb = load_workbook(path, data_only=True)
    ws = wb["Records"]
    lists = {}
    ls = wb["Lists"]
    for c in range(1, ls.max_column + 1, 3):
        name = ls.cell(1, c).value
        if not name:
            continue
        mapping = {}
        for r in range(2, ls.max_row + 1):
            label, value = ls.cell(r, c).value, ls.cell(r, c + 1).value
            if label is not None:
                mapping[str(label).strip()] = value
                mapping[str(value).strip()] = value
        lists[name] = mapping
    cols = {cdef["key"]: cdef for cdef in columns()}
    keys = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    notes: list[str] = []
    features = []
    info = {wb["File"].cell(r, 1).value: wb["File"].cell(r, 2).value
            for r in range(1, wb["File"].max_row + 1) if wb["File"].cell(r, 1).value}
    field_order = list(dict.fromkeys(c["field"] for c in columns())) + ["published-with-consent"]

    for r in range(FIRST_ROW, ws.max_row + 1):
        row = {k: _clean(ws.cell(r, c).value) for c, k in enumerate(keys, start=1) if k}
        if not row.get("id"):
            continue
        rid = _ref(str(row["id"]))
        props: dict = {}
        multi: dict[str, list] = {}
        contacts: dict[int, dict] = {}
        consent = []
        point = [None, None]
        for key, value in row.items():
            if value is None or key == "id" or key.startswith("review:"):
                continue
            cdef = cols.get(key)
            if cdef is None:
                if key.startswith("x-"):
                    yes_no = lists.get("yes-no", {}).get(str(value))
                    props[key] = (yes_no == "yes") if yes_no else value
                else:
                    notes.append(f"{rid}: unknown column {key} ignored")
                continue
            if cdef["list"]:
                mapped = lists[cdef["list"]].get(str(value))
                if mapped is None:
                    notes.append(f"{rid}: '{value}' is not in the {cdef['list']} list; ignored")
                    continue
                value = mapped
            if key.startswith("consent:"):
                if value == "yes":
                    consent.append(key.split(":", 1)[1])
            elif key.startswith("contacts."):
                _, n, part = key.split(".")
                contacts.setdefault(int(n), {})[part] = str(value)
            elif key in ("point.lon", "point.lat"):
                point[0 if key.endswith("lon") else 1] = float(value)
            elif key.endswith(".total"):
                props[key.split(".")[0]] = {"total": int(round(float(value)))}
            elif "." in key:
                field = key.split(".")[0]
                if field == "serves":
                    value = _ref(str(value))
                multi.setdefault(field, []).append(value)
            elif key == "parent":
                props["parent"] = _ref(str(value))
            elif key == "designated-area":
                props[key] = value == "yes"
            elif key in INT_FIELDS:
                props[key] = int(round(float(value)))
            elif key in NUM_FIELDS:
                number = float(value)
                props[key] = int(number) if number == int(number) else number
            elif key in DATE_FIELDS:
                props[key] = value.date().isoformat() if isinstance(value, dt.datetime) else (
                    value.isoformat() if isinstance(value, dt.date) else str(value))
            elif key in SPLIT_FIELDS:
                props[key] = [part.strip() for part in str(value).split(";") if part.strip()]
            elif key == "investors":
                props[key] = [{"name": part.strip()} for part in str(value).split(";") if part.strip()]
            else:
                props[key] = value if isinstance(value, (int, float)) else str(value)
        for field, values in multi.items():
            props[field] = list(dict.fromkeys(values))
        if contacts:
            props["contacts"] = [contacts[n] for n in sorted(contacts)]
        if consent:
            props["published-with-consent"] = consent
        if None not in point:
            props["point"] = [round(point[0], 6), round(point[1], 6)]
        # A value whose field does not apply to this kind is kept as a publisher field.
        kind = props.get("kind")
        for field in list(props):
            reg = REGISTRY.get(field)
            if reg and kind and reg["applies-to"] != "all" and kind not in reg["applies-to"].split():
                props["x-" + field] = props.pop(field)
                article = "an" if kind[0] in "aeiou" else "a"
                notes.append(f"{rid}: {field} does not apply to {article} {kind}; kept as x-{field}")
        ordered = {k: props[k] for k in field_order if k in props}
        ordered.update({k: v for k, v in props.items() if k not in ordered})
        geom = None
        if geometry and rid in geometry:
            geom = geometry[rid]
        elif "point" in props:
            geom = {"type": "Point", "coordinates": props["point"]}
        features.append({"type": "Feature", "id": rid, "geometry": geom, "properties": ordered})

    context = json.loads((ROOT / "schema" / "ilp-ds.context.jsonld").read_text(encoding="utf-8"))
    base = info.get("base")
    doc = {
        "@context": {**context["@context"], "@base": base},
        "type": "FeatureCollection",
        "standard": {"name": "ILP-DS", "version": _version()},
        "publisher": info.get("publisher"),
        "licence": info.get("licence"),
        "base": base,
        "changelog": [{"date": dt.date.today().isoformat(),
                       "note": info.get("change note") or "Converted from a review workbook"}],
        "features": features,
    }
    display = {}
    fs = wb["Fields"]
    if fs.max_column >= 7:
        for r in range(2, fs.max_row + 1):
            field, choice = fs.cell(r, 3).value, fs.cell(r, 7).value
            if field and choice in ("Yes", "No"):
                display[field] = choice == "Yes"
    return doc, notes, display


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = parser.add_subparsers(dest="command", required=True)
    rd = sub.add_parser("read", help="convert a reviewed workbook to ILP-DS")
    rd.add_argument("workbook", type=Path)
    rd.add_argument("--out", type=Path, required=True)
    rd.add_argument("--public", type=Path)
    rd.add_argument("--geometry", type=Path, help="GeoJSON of shapes, keyed by record id")
    rd.add_argument("--display-fields", type=Path, help="write the Fields sheet's display choices as JSON")
    args = parser.parse_args(argv)

    from check_ilp import check_document
    from export_public import make_public

    shapes = None
    if args.geometry:
        gj = json.loads(args.geometry.read_text(encoding="utf-8"))
        shapes = {_ref(f["id"]): f["geometry"] for f in gj["features"]}
    doc, notes, display = read_workbook(args.workbook, shapes)
    for note in notes:
        print("note:", note)
    problems = check_document(doc)
    args.out.write_text(json.dumps(doc, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {args.out} ({len(doc['features'])} records)")
    if problems:
        print("PROBLEMS in the internal file:")
        for problem in problems:
            print("  -", problem)
        return 1
    if args.public:
        public, removed = make_public(doc)
        public_problems = check_document(public)
        args.public.write_text(json.dumps(public, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"Wrote {args.public}; removed {len(removed)} non-public values")
        if public_problems:
            for problem in public_problems:
                print("  -", problem)
            return 1
    if args.display_fields:
        args.display_fields.write_text(json.dumps(display, indent=2) + "\n", encoding="utf-8")
    print("PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
